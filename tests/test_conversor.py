"""Teste do conversor DE VERDADE (LibreOffice + OCR), rodando no Docker com as restrições de produção.

    docker build -t agente-pu-conversor:teste conversor
    docker run -d --name conversor-teste --read-only --tmpfs /tmp:size=1g --cap-drop ALL \
        --security-opt no-new-privileges --memory 2g --pids-limit 256 -p 127.0.0.1:8097:8000 agente-pu-conversor:teste
    pip install python-pptx python-docx pypdfium2 pillow
    python tests/test_conversor.py

Gera os arquivos na hora (imagem com texto, PDF escaneado, PPTX/PPT com texto e print, DOCX com imagem) e confere
se o PDF devolvido tem tanto o texto digital quanto o texto que estava DENTRO das imagens (OCR).
"""
import io
import os
import subprocess
import sys
import tempfile

import httpx
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as PInches, Pt

URL = os.getenv("CONVERSOR_URL", "http://127.0.0.1:8097")
falhas = 0


def ok(cond, msg):
    global falhas
    print(("  OK  " if cond else "  FALHA ") + msg)
    falhas += 0 if cond else 1


def fonte(tamanho=46):
    for f in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(f, tamanho)
        except OSError:
            continue
    return ImageFont.load_default()


def imagem_com_texto(linhas, largura=1400) -> bytes:
    img = Image.new("RGB", (largura, 110 * len(linhas) + 60), "white")
    d = ImageDraw.Draw(img)
    for i, linha in enumerate(linhas):
        d.text((40, 40 + i * 110), linha, fill="black", font=fonte())
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def texto_pdf(pdf: bytes) -> str:
    doc = pdfium.PdfDocument(pdf)
    try:
        return "\n".join(p.get_textpage().get_text_bounded() for p in doc)
    finally:
        doc.close()


def converter(nome, dados):
    return httpx.post(f"{URL}/converter", params={"nome": nome}, content=dados, timeout=600)


ok(httpx.get(f"{URL}/saude", timeout=10).json() == {"ok": True}, "conversor no ar")

print("== imagem (foto/print de documento)")
png = imagem_com_texto(["FLEXX AG 2032", "Densidade 32 kg/m3", "Resiliencia 41 por cento"])
r = converter("ficha.png", png)
t = texto_pdf(r.content) if r.status_code == 200 else ""
ok(r.status_code == 200 and r.headers.get("x-ocr") == "aplicado", f"PNG vira PDF com OCR ({r.status_code}, {r.headers.get('x-ocr')})")
ok("FLEXX" in t and "2032" in t and "Densidade" in t, f"texto da imagem extraído: {t.strip()[:80]!r}")

print("== PDF escaneado (só imagem, sem texto digital)")
buf = io.BytesIO()
Image.open(io.BytesIO(imagem_com_texto(["Laudo de inspecao lote 24-1233", "Aprovado com desvio"]))).save(buf, "PDF", resolution=150)
r = converter("laudo-escaneado.pdf", buf.getvalue())
t = texto_pdf(r.content) if r.status_code == 200 else ""
ok(r.status_code == 200 and "24-1233" in t and "Aprovado" in t, f"OCR no PDF escaneado: {t.strip()[:80]!r}")

print("== PPTX com texto digital + print com texto")
with tempfile.TemporaryDirectory() as pasta:
    caminho_img = os.path.join(pasta, "print.png")
    with open(caminho_img, "wb") as f:
        f.write(imagem_com_texto(["Imagem: poliol vegetal 2026", "Fornecedor homologado"]))
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Compras de TDI 8020"
    caixa = s.shapes.add_textbox(PInches(0.5), PInches(1.5), PInches(9), PInches(1)).text_frame
    caixa.text = "Slide digital: consumo anual de TDI 8020"
    caixa.paragraphs[0].runs[0].font.size = Pt(24)
    s.shapes.add_picture(caminho_img, PInches(0.5), PInches(3), width=PInches(8))
    buf = io.BytesIO()
    p.save(buf)
    pptx = buf.getvalue()
    r = converter("apresentacao.pptx", pptx)
    t = texto_pdf(r.content) if r.status_code == 200 else ""
    ok(r.status_code == 200, f"PPTX convertido ({r.status_code}, {r.headers.get('x-paginas')} pág., OCR {r.headers.get('x-ocr')})")
    ok("TDI 8020" in t, "texto digital do slide preservado")
    ok("poliol" in t.lower() and "2026" in t, f"texto do print no slide extraído por OCR: {t.strip()[-90:]!r}")

    print("== PPT antigo (Office 97-2003)")
    with open(os.path.join(pasta, "apresentacao.pptx"), "wb") as f:
        f.write(pptx)
    gerou = subprocess.run(["docker", "run", "--rm", "-v", f"{pasta}:/w", "--entrypoint", "soffice", "agente-pu-conversor:teste",
                            "-env:UserInstallation=file:///tmp/p", "--headless", "--convert-to", "ppt", "--outdir", "/w",
                            "/w/apresentacao.pptx"], capture_output=True, timeout=300)
    caminho_ppt = os.path.join(pasta, "apresentacao.ppt")
    if os.path.exists(caminho_ppt):
        with open(caminho_ppt, "rb") as f:
            r = converter("apresentacao.ppt", f.read())
        t = texto_pdf(r.content) if r.status_code == 200 else ""
        ok(r.status_code == 200 and "TDI 8020" in t and "poliol" in t.lower(), "PPT antigo: texto digital + OCR do print")
    else:
        print("  (pulado: não consegui gerar um .ppt de teste)", gerou.stderr[-200:])

print("== DOCX com tabela/texto + imagem com texto")
d = Document()
d.add_heading("Especificacao FLEXX TH T190", 1)
d.add_paragraph("Texto digital: formulacao com poliol e TDI.")
d.add_picture(io.BytesIO(imagem_com_texto(["Grafico: custo por kg 18,40"])), width=Inches(5))
buf = io.BytesIO()
d.save(buf)
r = converter("especificacao.docx", buf.getvalue())
t = texto_pdf(r.content) if r.status_code == 200 else ""
ok(r.status_code == 200 and "formulacao" in t and "18,40" in t, "DOCX: texto digital + texto da imagem")
pdf_digital = r.content

print("== PDF digital (o texto existente é preservado)")
r = converter("digital.pdf", pdf_digital)
t = texto_pdf(r.content) if r.status_code == 200 else ""
ok(r.status_code == 200 and "formulacao" in t and r.headers.get("x-ocr") in ("aplicado", "parcial"), f"PDF digital reprocessado (OCR {r.headers.get('x-ocr')})")

print("== arquivos inválidos")
ok(converter("quebrado.pptx", b"isto nao e um pptx" * 50).status_code == 422, "PPTX corrompido -> 422")
ok(converter("script.exe", b"MZ" * 100).status_code == 415, "extensão não aceita -> 415")
ok(converter("vazio.pdf", b"").status_code == 400, "arquivo vazio -> 400")

print("\nfalhas:", falhas)
sys.exit(1 if falhas else 0)
